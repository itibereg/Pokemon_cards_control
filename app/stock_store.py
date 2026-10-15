from __future__ import annotations

import csv
import io
import re
import sqlite3
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable


DB_PATH = Path("data") / "stock.sqlite3"


class StockQuantityError(ValueError):
    """Erro de quantidade ao reservar/devolver unidades do estoque."""


@dataclass(slots=True)
class StockItem:
    id: int
    liga_card_id: str
    edition_id: str
    edition_code: str
    card_number: str
    edition_name: str
    rarity: str
    color: str
    name_pt: str
    name_en: str
    language: str
    quality: str
    quantity: int
    price: str
    extras: str
    source: str
    source_key: str
    updated_at: str

    @property
    def display_name(self) -> str:
        return self.name_pt or self.name_en or "Carta sem nome"

    @property
    def status(self) -> str:
        return "Disponível" if self.quantity > 0 else "Sem estoque"

    @staticmethod
    def _price_decimal(value: str) -> Decimal | None:
        raw = str(value or "").strip()
        if not raw:
            return None
        raw = raw.replace("R$", "").replace(" ", "")
        if "," in raw and "." in raw:
            raw = raw.replace(".", "").replace(",", ".")
        elif "," in raw:
            raw = raw.replace(",", ".")
        try:
            return Decimal(raw).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (InvalidOperation, ValueError):
            return None

    @property
    def unit_price_decimal(self) -> Decimal | None:
        return self._price_decimal(self.price)

    @property
    def price_display(self) -> str:
        value = self.unit_price_decimal
        if value is None:
            return ""
        integer, cents = f"{value:.2f}".split(".")
        integer = ".".join([integer[max(0, i - 3):i] for i in range(len(integer), 0, -3)][::-1])
        return f"R$ {integer},{cents}"

    @property
    def total_display(self) -> str:
        value = self.unit_price_decimal
        if value is None:
            return ""
        total = (value * self.quantity).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        integer, cents = f"{total:.2f}".split(".")
        integer = ".".join([integer[max(0, i - 3):i] for i in range(len(integer), 0, -3)][::-1])
        return f"R$ {integer},{cents}"


@dataclass(slots=True)
class StockChange:
    id: int
    stock_item_id: int
    action: str
    old_quantity: int
    new_quantity: int
    old_price: str
    new_price: str
    details: str
    created_at: str


@dataclass(slots=True)
class LigaImportResult:
    rows_read: int
    variants_imported: int
    inserted: int
    updated: int
    total_units: int
    skipped: int


class StockStore:
    """Persistência local do estoque e importação de arquivos da LigaPokemon."""

    VARIANT_COLUMNS = (
        ("Foil", "Foil"),
        ("Assinada", "Assinada"),
        ("Promo", "Promo"),
        ("Textless", "Textless"),
        ("Alterada", "Alterada"),
        ("Pre Release", "Pre Release"),
        ("Reverse Foil", "Reverse Foil"),
        ("Edition One", "1ª Edição"),
        ("Shadowless", "Shadowless"),
        ("Oversize", "Oversize"),
        ("Staff", "Staff"),
        ("Misprint", "Misprint"),
        ("Shattered Holo", "Shattered Holo"),
        ("Master Ball", "Master Ball"),
        ("Pokeball Foil", "Pokeball Foil"),
    )

    def __init__(self, db_path: Path | str = DB_PATH) -> None:
        self.db_path = Path(db_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self._initialize()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def _initialize(self) -> None:
        with self._connect() as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS stock_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    liga_card_id TEXT NOT NULL DEFAULT '',
                    edition_id TEXT NOT NULL DEFAULT '',
                    edition_code TEXT NOT NULL DEFAULT '',
                    card_number TEXT NOT NULL DEFAULT '',
                    edition_name TEXT NOT NULL DEFAULT '',
                    rarity TEXT NOT NULL DEFAULT '',
                    color TEXT NOT NULL DEFAULT '',
                    name_pt TEXT NOT NULL DEFAULT '',
                    name_en TEXT NOT NULL DEFAULT '',
                    language TEXT NOT NULL DEFAULT '',
                    quality TEXT NOT NULL DEFAULT '',
                    quantity INTEGER NOT NULL DEFAULT 0,
                    price TEXT NOT NULL DEFAULT '',
                    extras TEXT NOT NULL DEFAULT '',
                    source TEXT NOT NULL DEFAULT '',
                    source_key TEXT NOT NULL UNIQUE,
                    updated_at TEXT NOT NULL DEFAULT ''
                )
                """
            )
            conn.execute("CREATE INDEX IF NOT EXISTS idx_stock_name_pt ON stock_items(name_pt)")
            conn.execute("CREATE INDEX IF NOT EXISTS idx_stock_name_en ON stock_items(name_en)")
            conn.execute("CREATE INDEX IF NOT EXISTS idx_stock_number ON stock_items(card_number)")
            conn.execute("CREATE INDEX IF NOT EXISTS idx_stock_edition ON stock_items(edition_name)")
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS stock_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    stock_item_id INTEGER NOT NULL,
                    action TEXT NOT NULL DEFAULT '',
                    old_quantity INTEGER NOT NULL DEFAULT 0,
                    new_quantity INTEGER NOT NULL DEFAULT 0,
                    old_price TEXT NOT NULL DEFAULT '',
                    new_price TEXT NOT NULL DEFAULT '',
                    details TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT ''
                )
                """
            )
            conn.execute(
                "CREATE INDEX IF NOT EXISTS idx_stock_history_item "
                "ON stock_history(stock_item_id, id DESC)"
            )

    @staticmethod
    def _record_change(
        conn: sqlite3.Connection,
        *,
        item_id: int,
        action: str,
        old_quantity: int,
        new_quantity: int,
        old_price: str = "",
        new_price: str = "",
        details: str = "",
        created_at: str | None = None,
    ) -> None:
        conn.execute(
            """
            INSERT INTO stock_history (
                stock_item_id, action, old_quantity, new_quantity,
                old_price, new_price, details, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                int(item_id),
                str(action or ""),
                int(old_quantity),
                int(new_quantity),
                str(old_price or ""),
                str(new_price or ""),
                str(details or ""),
                created_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            ),
        )

    def get_history(self, item_id: int, limit: int = 30) -> list[StockChange]:
        with self._connect() as conn:
            rows = conn.execute(
                """
                SELECT *
                FROM stock_history
                WHERE stock_item_id = ?
                ORDER BY id DESC
                LIMIT ?
                """,
                (int(item_id), max(1, int(limit))),
            ).fetchall()
        return [StockChange(**dict(row)) for row in rows]

    def find_image_url(
        self,
        item: StockItem,
        cache_path: Path | str = Path("data") / "tcgdex_card_cache.sqlite3",
    ) -> str:
        """Tenta localizar a imagem da carta no cache local do TCGdex."""
        path = Path(cache_path)
        if not path.exists():
            return ""

        number = str(item.card_number or "").strip()
        name_pt = str(item.name_pt or "").strip()
        name_en = str(item.name_en or "").strip()
        if not number and not (name_pt or name_en):
            return ""

        try:
            conn = sqlite3.connect(path)
            conn.row_factory = sqlite3.Row
            try:
                row = conn.execute(
                    """
                    SELECT image_url
                    FROM tcgdex_cards
                    WHERE image_url <> ''
                      AND (
                        local_id = ? COLLATE NOCASE
                        OR number_text = ? COLLATE NOCASE
                        OR number_text LIKE ? COLLATE NOCASE
                      )
                      AND (
                        name = ? COLLATE NOCASE
                        OR name = ? COLLATE NOCASE
                        OR name LIKE ? COLLATE NOCASE
                        OR name LIKE ? COLLATE NOCASE
                      )
                    ORDER BY
                        CASE WHEN lang = 'pt' THEN 0 WHEN lang = 'en' THEN 1 ELSE 2 END,
                        rowid
                    LIMIT 1
                    """,
                    (
                        number,
                        number,
                        f"{number}/%",
                        name_pt,
                        name_en,
                        f"%{name_pt}%" if name_pt else "%",
                        f"%{name_en}%" if name_en else "%",
                    ),
                ).fetchone()
            finally:
                conn.close()
        except (sqlite3.Error, OSError):
            return ""

        return str(row["image_url"] or "") if row else ""

    @staticmethod
    def _normalize_header(value: str) -> str:
        text = str(value or "").replace("\n", " ").strip()
        text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
        return re.sub(r"\s+", " ", text).casefold()

    @classmethod
    def _header_lookup(cls, headers: Iterable[str]) -> dict[str, int]:
        return {cls._normalize_header(value): index for index, value in enumerate(headers)}

    @classmethod
    def _find_column(cls, lookup: dict[str, int], startswith: str) -> int | None:
        needle = cls._normalize_header(startswith)
        for key, index in lookup.items():
            if key == needle or key.startswith(needle):
                return index
        return None

    @staticmethod
    def _cell(row: list[str], index: int | None) -> str:
        if index is None or index >= len(row):
            return ""
        return str(row[index] or "").strip()

    @staticmethod
    def _parse_quantity(value: str) -> int:
        raw = str(value or "").strip().replace(",", ".")
        if not raw:
            return 0
        try:
            return max(0, int(float(raw)))
        except ValueError:
            return 0

    @staticmethod
    def _flag_enabled(value: str) -> bool:
        return str(value or "").strip().casefold() in {"1", "1.0", "sim", "s", "true", "yes", "x"}

    @classmethod
    def _build_extras(cls, row: list[str], lookup: dict[str, int]) -> str:
        extras: list[str] = []
        for header_prefix, label in cls.VARIANT_COLUMNS:
            idx = cls._find_column(lookup, header_prefix)
            if idx is not None and cls._flag_enabled(cls._cell(row, idx)):
                extras.append(label)
        return ", ".join(extras)

    @staticmethod
    def _source_key(*parts: str) -> str:
        clean = [str(part or "").strip().casefold() for part in parts]
        return "|".join(clean)

    @staticmethod
    def _decode_csv(path: Path) -> str:
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")

    @classmethod
    def _read_csv_rows(cls, path: Path) -> list[list[str]]:
        text = cls._decode_csv(path)
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]

    @classmethod
    def _read_xlsx_rows(cls, path: Path) -> list[list[str]]:
        # O projeto já utiliza openpyxl para exportação Excel. Aqui ele é usado
        # somente quando o usuário selecionar diretamente um .xlsx da LigaPokemon.
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Para importar .xlsx, instale openpyxl. Arquivos .csv funcionam normalmente.") from exc

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows: list[list[str]] = []
            for values in sheet.iter_rows(values_only=True):
                rows.append(["" if value is None else str(value) for value in values])
            return rows
        finally:
            workbook.close()

    @classmethod
    def _load_rows(cls, path: Path) -> list[list[str]]:
        suffix = path.suffix.lower()
        if suffix == ".csv":
            return cls._read_csv_rows(path)
        if suffix in {".xlsx", ".xlsm"}:
            return cls._read_xlsx_rows(path)
        raise ValueError("Formato não suportado. Selecione um arquivo .csv ou .xlsx exportado pela LigaPokemon.")

    @classmethod
    def _find_header_row(cls, rows: list[list[str]]) -> int:
        for index, row in enumerate(rows[:40]):
            normalized = {cls._normalize_header(cell) for cell in row}
            if "tipo" in normalized and "numero" in normalized and any(key.startswith("nome da carta pt") for key in normalized):
                return index
        raise ValueError("Não foi possível localizar o cabeçalho da planilha da LigaPokemon.")

    def import_ligapokemon(self, file_path: Path | str) -> LigaImportResult:
        """Importa/sincroniza variantes do arquivo exportado pela LigaPokemon.

        A quantidade do arquivo SUBSTITUI a quantidade da mesma variante já cadastrada.
        Isso permite importar o mesmo arquivo novamente sem duplicar o estoque.
        """
        path = Path(file_path)
        rows = self._load_rows(path)
        if not rows:
            raise ValueError("O arquivo selecionado está vazio.")

        header_index = self._find_header_row(rows)
        headers = rows[header_index]
        lookup = self._header_lookup(headers)

        col = {
            "edition_id": self._find_column(lookup, "Edição ID"),
            "edition_code": self._find_column(lookup, "Edição Sigla"),
            "card_id": self._find_column(lookup, "Carta ID"),
            "number": self._find_column(lookup, "Número"),
            "edition": self._find_column(lookup, "Edição"),
            "rarity": self._find_column(lookup, "Raridade"),
            "color": self._find_column(lookup, "Cor"),
            "name_pt": self._find_column(lookup, "Nome da Carta PT"),
            "name_en": self._find_column(lookup, "Nome da Carta EN"),
            "language": self._find_column(lookup, "Idioma"),
            "quality": self._find_column(lookup, "Qualidade"),
            "quantity_existing": self._find_column(lookup, "Quantidade Existente"),
            "quantity_change": self._find_column(lookup, "Quantidade Para Somar/Subtrair"),
            "price": self._find_column(lookup, "Preço"),
        }

        required = ("number", "edition", "name_pt", "name_en")
        if any(col[key] is None for key in required):
            raise ValueError("O arquivo não possui as colunas esperadas da exportação da LigaPokemon.")

        # Consolida linhas idênticas do mesmo arquivo. Normalmente haverá uma linha
        # por variante, mas isso evita perda caso o arquivo contenha duplicatas.
        imported: dict[str, dict[str, object]] = {}
        rows_read = 0
        skipped = 0

        for row in rows[header_index + 1 :]:
            if not any(str(cell or "").strip() for cell in row):
                continue
            rows_read += 1

            name_pt = self._cell(row, col["name_pt"])
            name_en = self._cell(row, col["name_en"])
            number = self._cell(row, col["number"])
            edition = self._cell(row, col["edition"])
            if not (name_pt or name_en) or not number:
                skipped += 1
                continue

            qty_change_raw = self._cell(row, col["quantity_change"])
            qty_existing_raw = self._cell(row, col["quantity_existing"])
            # Exportações de coleção da Liga costumam trazer a quantidade na coluna
            # "Quantidade Para Somar/Subtrair". Se ela estiver vazia, usa Existente.
            quantity = self._parse_quantity(qty_change_raw if qty_change_raw else qty_existing_raw)

            extras = self._build_extras(row, lookup)
            language = self._cell(row, col["language"]).upper() or "PT"
            quality = self._cell(row, col["quality"]).upper()
            card_id = self._cell(row, col["card_id"])
            edition_id = self._cell(row, col["edition_id"])
            edition_code = self._cell(row, col["edition_code"])

            key = self._source_key(
                "ligapokemon",
                card_id,
                edition_id,
                number,
                language,
                quality,
                extras,
                name_pt,
                name_en,
            )
            if key in imported:
                imported[key]["quantity"] = int(imported[key]["quantity"]) + quantity
            else:
                imported[key] = {
                    "liga_card_id": card_id,
                    "edition_id": edition_id,
                    "edition_code": edition_code,
                    "card_number": number,
                    "edition_name": edition,
                    "rarity": self._cell(row, col["rarity"]),
                    "color": self._cell(row, col["color"]),
                    "name_pt": name_pt,
                    "name_en": name_en,
                    "language": language,
                    "quality": quality,
                    "quantity": quantity,
                    "price": self._cell(row, col["price"]),
                    "extras": extras,
                    "source": "LigaPokemon",
                    "source_key": key,
                }

        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        inserted = 0
        updated = 0
        total_units = 0

        with self._connect() as conn:
            for item in imported.values():
                total_units += int(item["quantity"])
                existing = conn.execute(
                    "SELECT id, quantity, price FROM stock_items WHERE source_key = ?",
                    (item["source_key"],),
                ).fetchone()

                if existing:
                    old_quantity = int(existing["quantity"] or 0)
                    old_price = str(existing["price"] or "")
                    effective_price = str(item["price"] or "") or old_price
                    conn.execute(
                        """
                        UPDATE stock_items SET
                            liga_card_id=?, edition_id=?, edition_code=?, card_number=?, edition_name=?,
                            rarity=?, color=?, name_pt=?, name_en=?, language=?, quality=?, quantity=?,
                            price=CASE WHEN ? <> '' THEN ? ELSE price END, extras=?, source=?, updated_at=?
                        WHERE source_key=?
                        """,
                        (
                            item["liga_card_id"], item["edition_id"], item["edition_code"], item["card_number"],
                            item["edition_name"], item["rarity"], item["color"], item["name_pt"], item["name_en"],
                            item["language"], item["quality"], item["quantity"], item["price"], item["price"], item["extras"],
                            item["source"], now, item["source_key"],
                        ),
                    )
                    new_quantity = int(item["quantity"])
                    if old_quantity != new_quantity or old_price != effective_price:
                        self._record_change(
                            conn,
                            item_id=int(existing["id"]),
                            action="Sincronização Liga",
                            old_quantity=old_quantity,
                            new_quantity=new_quantity,
                            old_price=old_price,
                            new_price=effective_price,
                            details="Quantidade/preço sincronizados pela importação da LigaPokemon.",
                            created_at=now,
                        )
                    updated += 1
                else:
                    cursor = conn.execute(
                        """
                        INSERT INTO stock_items (
                            liga_card_id, edition_id, edition_code, card_number, edition_name,
                            rarity, color, name_pt, name_en, language, quality, quantity,
                            price, extras, source, source_key, updated_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            item["liga_card_id"], item["edition_id"], item["edition_code"], item["card_number"],
                            item["edition_name"], item["rarity"], item["color"], item["name_pt"], item["name_en"],
                            item["language"], item["quality"], item["quantity"], item["price"], item["extras"],
                            item["source"], item["source_key"], now,
                        ),
                    )
                    new_id = int(cursor.lastrowid)
                    self._record_change(
                        conn,
                        item_id=new_id,
                        action="Importação Liga",
                        old_quantity=0,
                        new_quantity=int(item["quantity"]),
                        old_price="",
                        new_price=str(item["price"] or ""),
                        details="Variante adicionada ao estoque pela LigaPokemon.",
                        created_at=now,
                    )
                    inserted += 1

        return LigaImportResult(
            rows_read=rows_read,
            variants_imported=len(imported),
            inserted=inserted,
            updated=updated,
            total_units=total_units,
            skipped=skipped,
        )

    def list_items(self, search: str = "", status: str = "Todos") -> list[StockItem]:
        query = "SELECT * FROM stock_items"
        where: list[str] = []
        params: list[object] = []

        search = str(search or "").strip()
        if search:
            like = f"%{search}%"
            where.append(
                "(name_pt LIKE ? COLLATE NOCASE OR name_en LIKE ? COLLATE NOCASE OR "
                "card_number LIKE ? COLLATE NOCASE OR edition_name LIKE ? COLLATE NOCASE OR "
                "edition_code LIKE ? COLLATE NOCASE OR extras LIKE ? COLLATE NOCASE OR language LIKE ? COLLATE NOCASE)"
            )
            params.extend([like] * 7)

        if status == "Disponível":
            where.append("quantity > 0")
        elif status == "Sem estoque":
            where.append("quantity <= 0")

        if where:
            query += " WHERE " + " AND ".join(where)
        query += " ORDER BY name_pt COLLATE NOCASE, name_en COLLATE NOCASE, edition_name COLLATE NOCASE, card_number COLLATE NOCASE"

        with self._connect() as conn:
            rows = conn.execute(query, params).fetchall()
        return [StockItem(**dict(row)) for row in rows]

    def search_available(self, search: str, limit: int = 12) -> list[StockItem]:
        """Busca somente variantes disponíveis para uso na tela de Vendas."""
        cleaned = " ".join(str(search or "").strip().split())
        if len(cleaned) < 2:
            return []

        like = f"%{cleaned}%"
        starts = f"{cleaned}%"
        with self._connect() as conn:
            rows = conn.execute(
                """
                SELECT *
                FROM stock_items
                WHERE quantity > 0
                  AND (
                    name_pt LIKE ? COLLATE NOCASE
                    OR name_en LIKE ? COLLATE NOCASE
                    OR card_number LIKE ? COLLATE NOCASE
                    OR edition_name LIKE ? COLLATE NOCASE
                    OR edition_code LIKE ? COLLATE NOCASE
                  )
                ORDER BY
                    CASE
                        WHEN name_pt LIKE ? COLLATE NOCASE THEN 0
                        WHEN name_en LIKE ? COLLATE NOCASE THEN 1
                        WHEN card_number LIKE ? COLLATE NOCASE THEN 2
                        ELSE 3
                    END,
                    name_pt COLLATE NOCASE,
                    name_en COLLATE NOCASE,
                    edition_name COLLATE NOCASE,
                    card_number COLLATE NOCASE,
                    language COLLATE NOCASE,
                    extras COLLATE NOCASE
                LIMIT ?
                """,
                (like, like, like, like, like, starts, starts, starts, int(limit)),
            ).fetchall()
        return [StockItem(**dict(row)) for row in rows]

    def get_item(self, item_id: int) -> StockItem | None:
        """Retorna uma variante específica do estoque."""
        with self._connect() as conn:
            row = conn.execute(
                "SELECT * FROM stock_items WHERE id = ?",
                (int(item_id),),
            ).fetchone()
        return StockItem(**dict(row)) if row is not None else None

    def reserve_quantity(self, item_id: int, quantity: int) -> StockItem:
        """Baixa quantidade de forma atômica e impede estoque negativo."""
        quantity = int(quantity)
        if quantity <= 0:
            raise StockQuantityError("A quantidade deve ser maior que zero.")

        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as conn:
            before = conn.execute(
                "SELECT quantity, price FROM stock_items WHERE id = ?",
                (int(item_id),),
            ).fetchone()
            cursor = conn.execute(
                """
                UPDATE stock_items
                SET quantity = quantity - ?, updated_at = ?
                WHERE id = ? AND quantity >= ?
                """,
                (quantity, now, int(item_id), quantity),
            )
            if cursor.rowcount != 1:
                row = conn.execute(
                    "SELECT quantity, name_pt, name_en FROM stock_items WHERE id = ?",
                    (int(item_id),),
                ).fetchone()
                if row is None:
                    raise StockQuantityError("A carta selecionada não existe mais no estoque.")
                available = int(row["quantity"] or 0)
                name = str(row["name_pt"] or row["name_en"] or "Carta")
                raise StockQuantityError(
                    f"Estoque insuficiente para {name}. Disponível: {available} unidade(s)."
                )

            old_q = int(before["quantity"] or 0) if before else quantity
            self._record_change(
                conn,
                item_id=int(item_id),
                action="Saída para venda",
                old_quantity=old_q,
                new_quantity=old_q - quantity,
                old_price=str(before["price"] or "") if before else "",
                new_price=str(before["price"] or "") if before else "",
                details=f"Reserva de {quantity} unidade(s) para venda.",
                created_at=now,
            )

        item = self.get_item(item_id)
        if item is None:
            raise StockQuantityError("A carta selecionada não existe mais no estoque.")
        return item

    def restore_quantity(self, item_id: int, quantity: int) -> None:
        """Devolve unidades ao estoque quando um item é retirado da venda."""
        quantity = int(quantity)
        if quantity <= 0:
            return
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as conn:
            before = conn.execute(
                "SELECT quantity, price FROM stock_items WHERE id = ?",
                (int(item_id),),
            ).fetchone()
            cursor = conn.execute(
                "UPDATE stock_items SET quantity = quantity + ?, updated_at = ? WHERE id = ?",
                (quantity, now, int(item_id)),
            )
            if cursor.rowcount != 1:
                raise StockQuantityError(
                    "Não foi possível devolver a carta ao estoque porque o item não existe mais."
                )
            old_q = int(before["quantity"] or 0) if before else 0
            price = str(before["price"] or "") if before else ""
            self._record_change(
                conn,
                item_id=int(item_id),
                action="Devolução",
                old_quantity=old_q,
                new_quantity=old_q + quantity,
                old_price=price,
                new_price=price,
                details=f"Devolução de {quantity} unidade(s) ao estoque.",
                created_at=now,
            )

    def replace_reservation(
        self,
        old_item_id: int,
        old_quantity: int,
        new_item_id: int,
        new_quantity: int,
    ) -> None:
        """Ajusta uma reserva já baixada sem permitir estoque negativo."""
        old_item_id = int(old_item_id)
        new_item_id = int(new_item_id)
        old_quantity = max(0, int(old_quantity))
        new_quantity = int(new_quantity)
        if new_quantity <= 0:
            raise StockQuantityError("A quantidade deve ser maior que zero.")

        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as conn:
            if old_item_id == new_item_id:
                row = conn.execute(
                    "SELECT quantity, price FROM stock_items WHERE id = ?",
                    (new_item_id,),
                ).fetchone()
                current_q = int(row["quantity"] or 0) if row else 0
                price = str(row["price"] or "") if row else ""
                delta = new_quantity - old_quantity

                if delta > 0:
                    cursor = conn.execute(
                        "UPDATE stock_items SET quantity = quantity - ?, updated_at = ? "
                        "WHERE id = ? AND quantity >= ?",
                        (delta, now, new_item_id, delta),
                    )
                    if cursor.rowcount != 1:
                        available = current_q
                        raise StockQuantityError(
                            f"Estoque insuficiente. Há {available} unidade(s) adicional(is) disponível(is)."
                        )
                    new_q = current_q - delta
                elif delta < 0:
                    conn.execute(
                        "UPDATE stock_items SET quantity = quantity + ?, updated_at = ? WHERE id = ?",
                        (-delta, now, old_item_id),
                    )
                    new_q = current_q + (-delta)
                else:
                    return

                self._record_change(
                    conn,
                    item_id=new_item_id,
                    action="Ajuste de venda",
                    old_quantity=current_q,
                    new_quantity=new_q,
                    old_price=price,
                    new_price=price,
                    details=f"Ajuste da reserva da venda: {old_quantity} → {new_quantity} unidade(s).",
                    created_at=now,
                )
                return

            old_row = conn.execute(
                "SELECT quantity, price FROM stock_items WHERE id = ?",
                (old_item_id,),
            ).fetchone()
            new_row = conn.execute(
                "SELECT quantity, price FROM stock_items WHERE id = ?",
                (new_item_id,),
            ).fetchone()
            if old_row is None:
                raise StockQuantityError("O item original da venda não existe mais no estoque.")
            if new_row is None:
                raise StockQuantityError("A nova carta selecionada não existe mais no estoque.")

            old_current = int(old_row["quantity"] or 0)
            new_current = int(new_row["quantity"] or 0)

            conn.execute(
                "UPDATE stock_items SET quantity = quantity + ?, updated_at = ? WHERE id = ?",
                (old_quantity, now, old_item_id),
            )
            cursor = conn.execute(
                "UPDATE stock_items SET quantity = quantity - ?, updated_at = ? "
                "WHERE id = ? AND quantity >= ?",
                (new_quantity, now, new_item_id, new_quantity),
            )
            if cursor.rowcount != 1:
                raise StockQuantityError(
                    f"Estoque insuficiente para a nova carta. Disponível: {new_current} unidade(s)."
                )

            self._record_change(
                conn,
                item_id=old_item_id,
                action="Troca na venda",
                old_quantity=old_current,
                new_quantity=old_current + old_quantity,
                old_price=str(old_row["price"] or ""),
                new_price=str(old_row["price"] or ""),
                details=f"{old_quantity} unidade(s) devolvida(s) após troca de variante.",
                created_at=now,
            )
            self._record_change(
                conn,
                item_id=new_item_id,
                action="Troca na venda",
                old_quantity=new_current,
                new_quantity=new_current - new_quantity,
                old_price=str(new_row["price"] or ""),
                new_price=str(new_row["price"] or ""),
                details=f"{new_quantity} unidade(s) reservada(s) na nova variante.",
                created_at=now,
            )

    def reserve_many(self, requests: Iterable[tuple[int, int]]) -> None:
        """Reserva várias variantes em uma única transação."""
        consolidated: dict[int, int] = {}
        for item_id, quantity in requests:
            quantity = int(quantity)
            if quantity <= 0:
                continue
            consolidated[int(item_id)] = consolidated.get(int(item_id), 0) + quantity
        if not consolidated:
            return

        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as conn:
            for item_id, quantity in consolidated.items():
                before = conn.execute(
                    "SELECT quantity, price FROM stock_items WHERE id = ?",
                    (item_id,),
                ).fetchone()
                available = int(before["quantity"] or 0) if before else 0
                cursor = conn.execute(
                    "UPDATE stock_items SET quantity = quantity - ?, updated_at = ? "
                    "WHERE id = ? AND quantity >= ?",
                    (quantity, now, item_id, quantity),
                )
                if cursor.rowcount != 1:
                    raise StockQuantityError(
                        f"Estoque insuficiente para duplicar a venda. "
                        f"Item {item_id}: disponível {available}, necessário {quantity}."
                    )
                price = str(before["price"] or "") if before else ""
                self._record_change(
                    conn,
                    item_id=item_id,
                    action="Saída para venda",
                    old_quantity=available,
                    new_quantity=available - quantity,
                    old_price=price,
                    new_price=price,
                    details=f"Reserva de {quantity} unidade(s) para venda.",
                    created_at=now,
                )

    def restore_many(self, requests: Iterable[tuple[int, int]]) -> None:
        """Devolve várias reservas ao estoque em uma única transação."""
        consolidated: dict[int, int] = {}
        for item_id, quantity in requests:
            quantity = int(quantity)
            if quantity <= 0:
                continue
            consolidated[int(item_id)] = consolidated.get(int(item_id), 0) + quantity
        if not consolidated:
            return

        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as conn:
            for item_id, quantity in consolidated.items():
                before = conn.execute(
                    "SELECT quantity, price FROM stock_items WHERE id = ?",
                    (item_id,),
                ).fetchone()
                if before is None:
                    continue
                old_q = int(before["quantity"] or 0)
                price = str(before["price"] or "")
                conn.execute(
                    "UPDATE stock_items SET quantity = quantity + ?, updated_at = ? WHERE id = ?",
                    (quantity, now, item_id),
                )
                self._record_change(
                    conn,
                    item_id=item_id,
                    action="Devolução",
                    old_quantity=old_q,
                    new_quantity=old_q + quantity,
                    old_price=price,
                    new_price=price,
                    details=f"Devolução de {quantity} unidade(s) ao estoque.",
                    created_at=now,
                )

    @staticmethod
    def normalize_price(value: str) -> str:
        raw = str(value or "").strip()
        if not raw:
            return ""
        raw = raw.replace("R$", "").replace(" ", "")
        if "," in raw and "." in raw:
            raw = raw.replace(".", "").replace(",", ".")
        elif "," in raw:
            raw = raw.replace(",", ".")
        try:
            amount = Decimal(raw).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (InvalidOperation, ValueError) as exc:
            raise ValueError("Preço inválido. Use por exemplo 12,50.") from exc
        if amount < 0:
            raise ValueError("O preço não pode ser negativo.")
        return f"{amount:.2f}"

    def update_price(self, item_id: int, price: str) -> None:
        normalized = self.normalize_price(price)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with self._connect() as conn:
            before = conn.execute(
                "SELECT quantity, price FROM stock_items WHERE id = ?",
                (int(item_id),),
            ).fetchone()
            if before is None:
                raise ValueError("A carta selecionada não existe mais no estoque.")

            old_price = str(before["price"] or "")
            quantity = int(before["quantity"] or 0)

            conn.execute(
                "UPDATE stock_items SET price = ?, updated_at = ? WHERE id = ?",
                (normalized, now, int(item_id)),
            )

            if old_price != normalized:
                self._record_change(
                    conn,
                    item_id=int(item_id),
                    action="Preço alterado",
                    old_quantity=quantity,
                    new_quantity=quantity,
                    old_price=old_price,
                    new_price=normalized,
                    details="Preço unitário atualizado manualmente.",
                    created_at=now,
                )

    def get_summary(self) -> tuple[int, int]:
        with self._connect() as conn:
            row = conn.execute(
                "SELECT COUNT(*) AS variants, COALESCE(SUM(quantity), 0) AS units FROM stock_items"
            ).fetchone()
        return int(row["variants"]), int(row["units"])
