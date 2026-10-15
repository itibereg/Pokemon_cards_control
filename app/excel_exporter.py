from __future__ import annotations
# VERSAO_ATUALIZADA: CUSTOMER_PHONE_V89_2026_09_28

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from app.models import (
    SaleItem,
    calculate_totals,
    ITEM_STATUS_SEPARATED,
    ITEM_STATUS_NOT_FOUND,
)


HEADERS = [
    "Carta",
    "Nome",
    "Nº Carta",
    "Nº Coleção",
    "Extras",
    "Idioma",
    "Situação",
    "Qtd",
    "Preço Unitário",
    "Total",
]


def export_sale_to_xlsx(
    items: list[SaleItem],
    output_path: str | Path,
    *,
    sale_title: str = "Venda Pokémon Cards",
    customer_name: str = "",
    customer_phone: str = "",
    status: str = "Aberta",
    notes: str = "",
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Venda"

    title = sale_title.strip() or "Venda Pokémon Cards"
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=16)
    sheet.merge_cells("A1:J1")

    metadata = [
        ("Cliente", customer_name or "-"),
        ("Telefone", customer_phone or "-"),
        ("Status", status or "-"),
        ("Emitido em", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    row = 2
    for label, value in metadata:
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=row, column=2, value=value)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        row += 1

    if notes.strip():
        sheet.cell(row=row, column=1, value="Observações").font = Font(bold=True)
        sheet.cell(row=row, column=2, value=notes.strip())
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        row += 1

    row += 1
    header_row = row
    for column_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    start_row = header_row + 1
    for item in items:
        sheet.append(
            [
                item.card.display_name,
                item.card.name or "-",
                item.card.card_number,
                item.card.collection_number,
                getattr(item, "extras", "") or "",
                getattr(item, "language", "PT") or "PT",
                getattr(item, "separation_status", "") or "Pendente",
                item.quantity,
                item.unit_price,
                item.net_total,
            ]
        )

    money_columns = [9, 10]
    for item_row in range(start_row, start_row + len(items)):
        for column in money_columns:
            sheet.cell(row=item_row, column=column).number_format = 'R$ #,##0.00'
        sheet.cell(row=item_row, column=6).alignment = Alignment(horizontal="center")
        sheet.cell(row=item_row, column=7).alignment = Alignment(horizontal="center")
        sheet.cell(row=item_row, column=8).alignment = Alignment(horizontal="center")

        status_value = str(sheet.cell(row=item_row, column=7).value or "")
        if status_value == ITEM_STATUS_SEPARATED:
            fill = PatternFill("solid", fgColor="E2F0D9")
            for column in range(1, 11):
                sheet.cell(row=item_row, column=column).fill = fill
        elif status_value == ITEM_STATUS_NOT_FOUND:
            fill = PatternFill("solid", fgColor="FCE4D6")
            for column in range(1, 11):
                sheet.cell(row=item_row, column=column).fill = fill

    totals = calculate_totals(items)
    total_row = start_row + len(items) + 1

    sheet.cell(row=total_row, column=8, value="Totais")
    sheet.cell(row=total_row, column=8).font = Font(bold=True)
    sheet.cell(row=total_row, column=9, value=totals.gross_total)
    sheet.cell(row=total_row, column=10, value=totals.net_total)

    for column in money_columns:
        sheet.cell(row=total_row, column=column).number_format = 'R$ #,##0.00'
        sheet.cell(row=total_row, column=column).font = Font(bold=True)

    widths = {
        "A": 38,
        "B": 24,
        "C": 12,
        "D": 12,
        "E": 18,
        "F": 10,
        "G": 18,
        "H": 8,
        "I": 16,
        "J": 16,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    thin = Side(style="thin", color="D9E2EC")
    for cells in sheet.iter_rows(min_row=header_row, max_row=total_row, min_col=1, max_col=10):
        for cell in cells:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(vertical="center")

    sheet.freeze_panes = f"A{start_row}"
    workbook.save(path)
    return path
